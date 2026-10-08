"""
CrowdStrike / Inventory hostname matcher.

Expected layout (row 1 = headers):
    Column A: Hostname (from CrowdStrike)
    Column B: Last Name (inventory)
    Column C: Name      (inventory)
    Column D: ID#       (inventory, 5-character HSI)

A hostname matches when its last 5 characters equal an ID# in column D
(case-insensitive, surrounding spaces ignored).

Usage:
    python match_hostnames.py data.xlsx                 # adds/replaces a "Matched" tab in data.xlsx
    python match_hostnames.py data.xlsx -o result.xlsx  # writes to a new file instead
    python match_hostnames.py data.xlsx --sheet Sheet1  # pick the source tab (default: first tab)
    python match_hostnames.py data.csv                  # writes data_matched.csv

CSV output has a Status column on every row:
    Matched                      - hostname whose last 5 chars match an ID#
    Unmatched - hostname only    - hostname with no matching ID#
    Unmatched - inventory only   - inventory item no hostname matched

Requires: openpyxl  (pip install openpyxl)
"""

import argparse
import csv
import sys
from pathlib import Path

MATCHED_TAB = "Matched"
HEADERS = ["Hostname", "ID#", "Last Name", "Name"]


def clean(value):
    """Turn a cell value into a trimmed string ('' for empty)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # 84243.0 -> "84243"
    return str(value).strip()


def find_matches(rows):
    """rows: list of [A, B, C, D] lists, header excluded. Returns (matches, warnings)."""
    inventory = {}
    short_ids = []
    for r in rows:
        last, name, hsi = clean(r[1]), clean(r[2]), clean(r[3])
        if not hsi:
            continue
        if len(hsi) != 5:
            short_ids.append(hsi)
        inventory.setdefault(hsi.upper(), []).append((hsi, last, name))

    matches = []
    shared = {}
    for r in rows:
        host = clean(r[0])
        if not host:
            continue
        key = host[-5:].upper()
        for hsi, last, name in inventory.get(key, []):
            matches.append([host, hsi, last, name])
            shared.setdefault(key, set()).add(host)

    warnings = []
    if short_ids:
        warnings.append(
            f"{len(short_ids)} ID# value(s) in column D are not 5 characters and can't match "
            f"(possible lost leading zeros), e.g. {', '.join(short_ids[:5])}"
        )
    dup_inv = [k for k, v in inventory.items() if len(v) > 1 and k in shared]
    if dup_inv:
        warnings.append(f"ID#s listed more than once in inventory: {', '.join(dup_inv)}")
    multi_host = {k: sorted(v) for k, v in shared.items() if len(v) > 1}
    for k, hosts in multi_host.items():
        warnings.append(f"ID# {k} matched by multiple hostnames: {', '.join(hosts)}")
    return matches, warnings


STATUS_MATCHED = "Matched"
STATUS_HOST_ONLY = "Unmatched - hostname only"
STATUS_INV_ONLY = "Unmatched - inventory only"
FULL_HEADERS = ["Status"] + HEADERS


def full_report(rows, matches):
    """Matched rows first, then hostnames with no inventory match,
    then inventory items no hostname matched."""
    matched_hosts = {m[0] for m in matches}
    matched_ids = {m[1].upper() for m in matches}

    out = [[STATUS_MATCHED] + m for m in matches]
    for r in rows:
        host = clean(r[0])
        if host and host not in matched_hosts:
            out.append([STATUS_HOST_ONLY, host, "", "", ""])
    for r in rows:
        last, name, hsi = clean(r[1]), clean(r[2]), clean(r[3])
        if (hsi or last or name) and hsi.upper() not in matched_ids:
            out.append([STATUS_INV_ONLY, "", hsi, last, name])
    return out


def pad(row, n=4):
    row = list(row)
    return (row + [None] * n)[:n]


def run_xlsx(path, sheet_name, out_path):
    from openpyxl import load_workbook
    from openpyxl.styles import Font

    keep_vba = path.suffix.lower() == ".xlsm"
    wb = load_workbook(path, data_only=True, keep_vba=keep_vba)
    if sheet_name:
        if sheet_name not in wb.sheetnames:
            sys.exit(f"Tab '{sheet_name}' not found. Tabs: {', '.join(wb.sheetnames)}")
        ws = wb[sheet_name]
    else:
        ws = next(s for s in wb.worksheets if s.title != MATCHED_TAB)

    all_rows = [pad(r) for r in ws.iter_rows(min_col=1, max_col=4, values_only=True)]
    if not all_rows:
        sys.exit(f"Tab '{ws.title}' is empty.")
    check_header(all_rows[0], ws.title)
    matches, warnings = find_matches(all_rows[1:])

    if MATCHED_TAB in wb.sheetnames:
        del wb[MATCHED_TAB]
    out = wb.create_sheet(MATCHED_TAB)
    out.append(HEADERS)
    for m in matches:
        out.append(m)
    for cell in out[1]:
        cell.font = Font(bold=True)
    out.freeze_panes = "A2"
    for col, width in zip("ABCD", (22, 10, 26, 30)):
        out.column_dimensions[col].width = width
    # Keep IDs like 01XLK / 84243 as text so leading zeros survive
    for (cell,) in out.iter_rows(min_row=2, min_col=2, max_col=2):
        cell.number_format = "@"

    wb.save(out_path)
    return ws.title, matches, warnings


def run_csv(path, out_path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        all_rows = [pad(r) for r in csv.reader(f)]
    if not all_rows:
        sys.exit("CSV is empty.")
    check_header(all_rows[0], path.name)
    data = all_rows[1:]
    matches, warnings = find_matches(data)
    report = full_report(data, matches)
    with open(out_path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(FULL_HEADERS)
        w.writerows(report)
    host_only = sum(1 for r in report if r[0] == STATUS_HOST_ONLY)
    inv_only = sum(1 for r in report if r[0] == STATUS_INV_ONLY)
    warnings.insert(0, f"Also wrote {host_only} unmatched hostname(s) and "
                       f"{inv_only} unmatched inventory item(s).")
    return path.name, matches, warnings


def check_header(header, label):
    if not clean(header[0]) or not clean(header[3]):
        print(
            f"Warning: row 1 of '{label}' doesn't look like the expected layout "
            "(Hostname in A, ID# in D). Continuing anyway.",
            file=sys.stderr,
        )


def main():
    p = argparse.ArgumentParser(description="Match CrowdStrike hostnames to 5-character inventory HSIs.")
    p.add_argument("file", type=Path, help=".xlsx, .xlsm or .csv file")
    p.add_argument("--sheet", help="source tab name (xlsx only; default: first tab)")
    p.add_argument("-o", "--output", type=Path, help="output file (default: xlsx is updated in place; csv gets _matched.csv)")
    args = p.parse_args()

    if not args.file.exists():
        sys.exit(f"File not found: {args.file}")
    ext = args.file.suffix.lower()

    if ext in (".xlsx", ".xlsm"):
        out_path = args.output or args.file
        source, matches, warnings = run_xlsx(args.file, args.sheet, out_path)
        where = f"'{MATCHED_TAB}' tab in {out_path}"
    elif ext == ".csv":
        out_path = args.output or args.file.with_name(args.file.stem + "_matched.csv")
        source, matches, warnings = run_csv(args.file, out_path)
        where = str(out_path)
    else:
        sys.exit("Unsupported file type. Use .xlsx, .xlsm or .csv.")

    print(f"Source: {source}")
    print(f"Found {len(matches)} matching hostname(s). Written to {where}.")
    for w in warnings:
        print(f"Note: {w}")


if __name__ == "__main__":
    main()